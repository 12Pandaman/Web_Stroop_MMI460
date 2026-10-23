import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_research_template():
    wb = openpyxl.Workbook()

    # -------------------------------------------------------------
    # SHEET 1: Master Pre vs Post Comparison
    # -------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "สรุปผล Pre_vs_Post"

    # Title Banner
    ws1.merge_cells('A1:I1')
    ws1['A1'] = "ตารางสรุปผลการวิจัย Stroop Effect เปรียบเทียบ Pre-test และ Post-test"
    ws1['A1'].font = Font(name='Segoe UI', size=14, bold=True, color='FFFFFF')
    ws1['A1'].fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
    ws1['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws1.row_dimensions[1].height = 36

    # Subtitle
    ws1.merge_cells('A2:I2')
    ws1['A2'] = "เวลาปฏิกิริยาเฉลี่ย (Reaction Time: ms) เฉพาะข้อที่ตอบถูก | ผลต่าง D = Incongruent - Congruent | Delta D = Post - Pre"
    ws1['A2'].font = Font(name='Segoe UI', size=9, italic=True, color='64748B')
    ws1['A2'].alignment = Alignment(horizontal='center', vertical='center')
    ws1.row_dimensions[2].height = 20

    # Header Row 1 (Row 4)
    ws1.merge_cells('A4:A5')
    ws1['A4'] = "ผู้เข้าร่วม"

    ws1.merge_cells('B4:D4')
    ws1['B4'] = "Pre-test (ก่อนการทดลอง)"

    ws1.merge_cells('E4:G4')
    ws1['E4'] = "Post-test (หลังการทดลอง)"

    ws1.merge_cells('H4:H5')
    ws1['H4'] = "ผลต่าง Delta D\n(Post - Pre)"

    ws1.merge_cells('I4:I5')
    ws1['I4'] = "การเปลี่ยนแปลง\n(Stroop Effect)"

    # Header Row 2 (Row 5)
    ws1['B5'] = "Congruent (ms)"
    ws1['C5'] = "Incongruent (ms)"
    ws1['D5'] = "ผลต่าง D_pre"

    ws1['E5'] = "Congruent (ms)"
    ws1['F5'] = "Incongruent (ms)"
    ws1['G5'] = "ผลต่าง D_post"

    fill_main = PatternFill(start_color='334155', end_color='334155', fill_type='solid')
    fill_pre = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
    fill_post = PatternFill(start_color='065F46', end_color='065F46', fill_type='solid')
    fill_diff = PatternFill(start_color='4C1D95', end_color='4C1D95', fill_type='solid')
    font_head = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')
    border_thin = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    for cell in ['A4', 'A5']:
        ws1[cell].fill = fill_main
        ws1[cell].font = font_head
        ws1[cell].alignment = Alignment(horizontal='center', vertical='center')

    for col in ['B4', 'C4', 'D4', 'B5', 'C5', 'D5']:
        ws1[col].fill = fill_pre
        ws1[col].font = font_head
        ws1[col].alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for col in ['E4', 'F4', 'G4', 'E5', 'F5', 'G5']:
        ws1[col].fill = fill_post
        ws1[col].font = font_head
        ws1[col].alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for cell in ['H4', 'H5', 'I4', 'I5']:
        ws1[cell].fill = fill_diff
        ws1[cell].font = font_head
        ws1[cell].alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    ws1.row_dimensions[4].height = 26
    ws1.row_dimensions[5].height = 26

    # Sample participants matching the user's uploaded image
    sample_data = [
        ('P01', 500, 650, 480, 580),
        ('P02', 550, 700, 520, 630),
        ('P03', 480, 600, 460, 550),
        ('P04', 520, 680, 500, 610),
        ('P05', 510, 660, 490, 590),
    ]

    for idx, (p, pre_c, pre_i, post_c, post_i) in enumerate(sample_data, start=6):
        ws1[f'A{idx}'] = p
        ws1[f'B{idx}'] = pre_c
        ws1[f'C{idx}'] = pre_i
        ws1[f'D{idx}'] = f'=C{idx}-B{idx}' # D_pre = I - C
        ws1[f'E{idx}'] = post_c
        ws1[f'F{idx}'] = post_i
        ws1[f'G{idx}'] = f'=F{idx}-E{idx}' # D_post = I - C
        ws1[f'H{idx}'] = f'=G{idx}-D{idx}' # Delta D = D_post - D_pre
        ws1[f'I{idx}'] = f'=IF(H{idx}<0, "ลดลง " & ABS(H{idx}) & " ms (ดีขึ้น)", IF(H{idx}>0, "เพิ่มขึ้น " & H{idx} & " ms", "เท่าเดิม"))'

        ws1[f'A{idx}'].alignment = Alignment(horizontal='center', vertical='center')
        ws1[f'A{idx}'].font = Font(name='Segoe UI', size=10, bold=True)

        for col_letter in ['B', 'C', 'D', 'E', 'F', 'G', 'H']:
            ws1[f'{col_letter}{idx}'].alignment = Alignment(horizontal='right', vertical='center')
            ws1[f'{col_letter}{idx}'].font = Font(name='Segoe UI', size=10)

        ws1[f'I{idx}'].alignment = Alignment(horizontal='center', vertical='center')
        ws1[f'I{idx}'].font = Font(name='Segoe UI', size=10)

        for col_letter in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I']:
            ws1[f'{col_letter}{idx}'].border = border_thin

        ws1.row_dimensions[idx].height = 22

    # Average / Mean Row
    avg_row = len(sample_data) + 6
    ws1[f'A{avg_row}'] = "ค่าเฉลี่ย (Mean)"
    ws1[f'A{avg_row}'].font = Font(name='Segoe UI', size=10, bold=True)
    ws1[f'A{avg_row}'].alignment = Alignment(horizontal='center', vertical='center')
    fill_avg = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')

    for col_letter in ['B', 'C', 'D', 'E', 'F', 'G', 'H']:
        ws1[f'{col_letter}{avg_row}'] = f'=ROUND(AVERAGE({col_letter}6:{col_letter}{avg_row-1}), 1)'
        ws1[f'{col_letter}{avg_row}'].font = Font(name='Segoe UI', size=10, bold=True)
        ws1[f'{col_letter}{avg_row}'].alignment = Alignment(horizontal='right', vertical='center')

    ws1[f'I{avg_row}'] = f'=IF(H{avg_row}<0, "ลดลงเฉลี่ย " & ABS(H{avg_row}) & " ms (ดีขึ้น)", "-")'
    ws1[f'I{avg_row}'].font = Font(name='Segoe UI', size=10, bold=True)
    ws1[f'I{avg_row}'].alignment = Alignment(horizontal='center', vertical='center')

    for col_letter in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I']:
        ws1[f'{col_letter}{avg_row}'].fill = fill_avg
        ws1[f'{col_letter}{avg_row}'].border = border_thin
    ws1.row_dimensions[avg_row].height = 25

    # -------------------------------------------------------------
    # SHEET 2: Individual Tables matching exact user photo (Pre & Post)
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="ตารางเดี่ยว_Pre_และ_Post")

    # Table A: Pre-test
    ws2['A1'] = "ตารางผลการทดสอบ: Pre-test (ก่อนการทดลอง)"
    ws2['A1'].font = Font(name='Segoe UI', size=12, bold=True, color='1E3A8A')
    ws2.merge_cells('A1:D1')

    ws2['A3'] = "ผู้เข้าร่วม"
    ws2['B3'] = "Congruent (ms)"
    ws2['C3'] = "Incongruent (ms)"
    ws2['D3'] = "ผลต่าง D = I - C"

    for col in ['A3', 'B3', 'C3', 'D3']:
        ws2[col].fill = fill_pre
        ws2[col].font = font_head
        ws2[col].alignment = Alignment(horizontal='center', vertical='center')
        ws2[col].border = border_thin
    ws2.row_dimensions[3].height = 25

    for idx, (p, pre_c, pre_i, _, _) in enumerate(sample_data, start=4):
        ws2[f'A{idx}'] = p
        ws2[f'B{idx}'] = pre_c
        ws2[f'C{idx}'] = pre_i
        ws2[f'D{idx}'] = f'=C{idx}-B{idx}'

        ws2[f'A{idx}'].alignment = Alignment(horizontal='center', vertical='center')
        ws2[f'B{idx}'].alignment = Alignment(horizontal='right', vertical='center')
        ws2[f'C{idx}'].alignment = Alignment(horizontal='right', vertical='center')
        ws2[f'D{idx}'].alignment = Alignment(horizontal='right', vertical='center')

        for col in ['A', 'B', 'C', 'D']:
            ws2[f'{col}{idx}'].font = Font(name='Segoe UI', size=10)
            ws2[f'{col}{idx}'].border = border_thin
        ws2.row_dimensions[idx].height = 22

    pre_avg = len(sample_data) + 4
    ws2[f'A{pre_avg}'] = "เฉลี่ย"
    ws2[f'B{pre_avg}'] = f'=ROUND(AVERAGE(B4:B{pre_avg-1}), 1)'
    ws2[f'C{pre_avg}'] = f'=ROUND(AVERAGE(C4:C{pre_avg-1}), 1)'
    ws2[f'D{pre_avg}'] = f'=C{pre_avg}-B{pre_avg}'
    for col in ['A', 'B', 'C', 'D']:
        ws2[f'{col}{pre_avg}'].fill = fill_avg
        ws2[f'{col}{pre_avg}'].font = Font(name='Segoe UI', size=10, bold=True)
        ws2[f'{col}{pre_avg}'].border = border_thin
        if col != 'A':
            ws2[f'{col}{pre_avg}'].alignment = Alignment(horizontal='right', vertical='center')
        else:
            ws2[f'{col}{pre_avg}'].alignment = Alignment(horizontal='center', vertical='center')

    # Table B: Post-test
    post_start = pre_avg + 3
    ws2[f'A{post_start}'] = "ตารางผลการทดสอบ: Post-test (หลังการทดลอง)"
    ws2[f'A{post_start}'].font = Font(name='Segoe UI', size=12, bold=True, color='065F46')
    ws2.merge_cells(f'A{post_start}:D{post_start}')

    h_row = post_start + 2
    ws2[f'A{h_row}'] = "ผู้เข้าร่วม"
    ws2[f'B{h_row}'] = "Congruent (ms)"
    ws2[f'C{h_row}'] = "Incongruent (ms)"
    ws2[f'D{h_row}'] = "ผลต่าง D = I - C"

    for col in ['A', 'B', 'C', 'D']:
        ws2[f'{col}{h_row}'].fill = fill_post
        ws2[f'{col}{h_row}'].font = font_head
        ws2[f'{col}{h_row}'].alignment = Alignment(horizontal='center', vertical='center')
        ws2[f'{col}{h_row}'].border = border_thin
    ws2.row_dimensions[h_row].height = 25

    for idx, (p, _, _, post_c, post_i) in enumerate(sample_data, start=h_row+1):
        ws2[f'A{idx}'] = p
        ws2[f'B{idx}'] = post_c
        ws2[f'C{idx}'] = post_i
        ws2[f'D{idx}'] = f'=C{idx}-B{idx}'

        ws2[f'A{idx}'].alignment = Alignment(horizontal='center', vertical='center')
        ws2[f'B{idx}'].alignment = Alignment(horizontal='right', vertical='center')
        ws2[f'C{idx}'].alignment = Alignment(horizontal='right', vertical='center')
        ws2[f'D{idx}'].alignment = Alignment(horizontal='right', vertical='center')

        for col in ['A', 'B', 'C', 'D']:
            ws2[f'{col}{idx}'].font = Font(name='Segoe UI', size=10)
            ws2[f'{col}{idx}'].border = border_thin
        ws2.row_dimensions[idx].height = 22

    post_avg = h_row + len(sample_data) + 1
    ws2[f'A{post_avg}'] = "เฉลี่ย"
    ws2[f'B{post_avg}'] = f'=ROUND(AVERAGE(B{h_row+1}:B{post_avg-1}), 1)'
    ws2[f'C{post_avg}'] = f'=ROUND(AVERAGE(C{h_row+1}:C{post_avg-1}), 1)'
    ws2[f'D{post_avg}'] = f'=C{post_avg}-B{post_avg}'
    for col in ['A', 'B', 'C', 'D']:
        ws2[f'{col}{post_avg}'].fill = fill_avg
        ws2[f'{col}{post_avg}'].font = Font(name='Segoe UI', size=10, bold=True)
        ws2[f'{col}{post_avg}'].border = border_thin
        if col != 'A':
            ws2[f'{col}{post_avg}'].alignment = Alignment(horizontal='right', vertical='center')
        else:
            ws2[f'{col}{post_avg}'].alignment = Alignment(horizontal='center', vertical='center')

    # Adjust column widths
    for ws in [ws1, ws2]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                if '\n' in val_str:
                    val_str = max(val_str.split('\n'), key=len)
                max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 16)

    output_path = 'Stroop_PrePost_Research_Summary.xlsx'
    wb.save(output_path)
    print(f'Created {output_path} successfully!')

if __name__ == '__main__':
    build_research_template()
